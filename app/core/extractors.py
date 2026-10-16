from app.config import TEXT_EXTENSIONS, MAX_TEXT_BYTES

def extract_text(path):
    suffix = path.suffix.lower()
    try:
        if path.stat().st_size > MAX_TEXT_BYTES:
            return ""
        if suffix in {".pdf"}:
            try:
                from pypdf import PdfReader
                reader = PdfReader(str(path))
                parts = []
                for page in reader.pages[:40]:
                    parts.append(page.extract_text() or "")
                return "\\n".join(parts)[:500_000]
            except Exception:
                return ""
        if suffix == ".docx":
            try:
                from docx import Document
                doc = Document(str(path))
                parts = [p.text for p in doc.paragraphs if p.text]
                for table in doc.tables[:20]:
                    for row in table.rows[:100]:
                        parts.append(" | ".join(cell.text for cell in row.cells))
                return "\\n".join(parts)[:500_000]
            except Exception:
                return ""
        if suffix not in TEXT_EXTENSIONS:
            return ""
        raw = path.read_bytes()
        if b"\\x00" in raw[:8192]:
            return ""
        return raw.decode("utf-8", errors="replace")
    except (OSError, ValueError):
        return ""
